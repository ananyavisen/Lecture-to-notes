from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import time

def create_docx(notes, output_file):
    start_time = time.perf_counter()
    document = Document()

    # -----------------------------
    # Page setup
    # -----------------------------

    section = document.sections[0]

    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    # -----------------------------
    # Default font
    # -----------------------------

    normal_style = document.styles["Normal"]
    normal_style.font.name = "Calibri"
    normal_style.font.size = Pt(10.5)

    # -----------------------------
    # Title
    # -----------------------------

    title = document.add_paragraph()

    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = title.add_run(notes["title"])

    run.bold = True
    run.font.size = Pt(18)

    # -----------------------------
    # Major sections
    # -----------------------------

    for section_data in notes.get("sections", []):

        number = section_data.get("number", "")
        section_title = section_data.get("title", "")

        heading = document.add_paragraph()

        heading.paragraph_format.space_before = Pt(8)
        heading.paragraph_format.space_after = Pt(3)

        run = heading.add_run(
            f"{number}. {section_title}"
        )

        run.bold = True
        run.font.size = Pt(13)

        # -----------------------------
        # Subsections
        # -----------------------------

        for subsection in section_data.get("subsections", []):

            subsection_title = subsection.get("title", "")
            content = subsection.get("content", "")

            # Subsection heading
            subheading = document.add_paragraph()

            subheading.paragraph_format.space_before = Pt(4)
            subheading.paragraph_format.space_after = Pt(1)

            run = subheading.add_run(subsection_title)

            run.bold = True
            run.font.size = Pt(10.5)

            # Subsection content
            paragraph = document.add_paragraph()

            paragraph.paragraph_format.space_after = Pt(3)
            paragraph.paragraph_format.line_spacing = 1.05

            paragraph.add_run(content)

    # -----------------------------
    # Workflows
    # -----------------------------

    workflows = notes.get("workflows", [])

    if workflows:

        heading = document.add_paragraph()

        run = heading.add_run("Workflows")

        run.bold = True
        run.font.size = Pt(13)

        for workflow in workflows:

            paragraph = document.add_paragraph()

            paragraph.style = document.styles["List Bullet"]

            paragraph.add_run(workflow)

    # -----------------------------
    # Important takeaways
    # -----------------------------

    takeaways = notes.get("important_takeaways", [])

    if takeaways:

        heading = document.add_paragraph()

        run = heading.add_run("Important Takeaways")

        run.bold = True
        run.font.size = Pt(13)

        for takeaway in takeaways:

            paragraph = document.add_paragraph()

            paragraph.style = document.styles["List Bullet"]

            paragraph.add_run(takeaway)

    # -----------------------------
    # Common mistakes
    # -----------------------------

    mistakes = notes.get("common_mistakes", [])

    if mistakes:

        heading = document.add_paragraph()

        run = heading.add_run("Common Beginner Mistakes")

        run.bold = True
        run.font.size = Pt(13)

        for mistake in mistakes:

            paragraph = document.add_paragraph()

            paragraph.style = document.styles["List Bullet"]

            paragraph.add_run(mistake)

    # -----------------------------
    # Save
    # -----------------------------

    document.save(output_file)
    end_time = time.perf_counter()

    print(
        f"DOCX generation time: "
        f"{end_time - start_time:.4f} seconds"
    )
    print(f"Smart notes saved to: {output_file}")